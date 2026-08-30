#!/usr/bin/env python3
"""Pomiar odniesień do rycin i wyróżników zdradzających odpowiedź. Instrument akceptacyjny planu detells."""
import collections
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from generate_quiz_data import cell_text, is_excluded  # noqa: E402

LETTER_TO_INDEX = {"A": 0, "B": 1, "C": 2, "D": 3}
IMG = re.compile(r"\bimage\b", re.IGNORECASE)
IMG_PAREN = re.compile(r"\s*\((?:zob\.\s*|por\.\s*)?Image[^)]*\)", re.IGNORECASE)
CONJ = re.compile(r"\S\s+(i|oraz)\s+\S", re.IGNORECASE)
PAREN = re.compile(r"\([^)]{3,}\)")
COMMA = re.compile(r",")
CATG = re.compile(r"\b(tylko|wyłącznie|nigdy|zawsze|jedynie|każd\w+|wszystk\w+|żad\w+)\b", re.IGNORECASE)


def load_questions(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    out = []
    for sheet in [s for s in wb.sheetnames if s != "Struktura"]:
        ws = wb[sheet]
        for row in range(2, ws.max_row + 1):
            vals = [cell_text(ws.cell(row, c)) for c in range(1, 11)]
            qid, q, a, b, c_, d, correct, expl, _tags, excl = vals
            if q == "Pytanie" and a == "A":
                continue
            if q and not any([a, b, c_, d]):
                continue
            if is_excluded(excl):
                continue
            if not all([q, a, b, c_, d]):
                continue
            ci = LETTER_TO_INDEX.get(correct.strip().upper()[:1] if correct else "")
            if ci is None:
                continue
            out.append({"sheet": sheet, "row": row, "id": qid, "question": q, "options": [a, b, c_, d], "correct": ci, "explanation": expl})
    return out


def only_correct(qs, pat):
    hits = []
    for q in qs:
        flags = [bool(pat.search(o)) for o in q["options"]]
        if flags[q["correct"]] and sum(flags) == 1:
            hits.append(q)
    return hits


def main():
    if len(sys.argv) != 3:
        sys.exit("Użycie: measure_tells.py <xlsx> <out.json>")
    qs = load_questions(sys.argv[1])
    n = len(qs)
    img_qs = [q for q in qs if IMG.search(q["explanation"])]
    par = sum(1 for q in img_qs if not IMG.search(IMG_PAREN.sub("", q["explanation"])))
    catg_wrong, len15, longest, shortest, strong_shortest = [], [], 0, 0, 0
    tok_c, tok_w = collections.Counter(), collections.Counter()
    for q in qs:
        flags = [bool(CATG.search(o)) for o in q["options"]]
        if not flags[q["correct"]] and sum(flags) >= 2:
            catg_wrong.append(q)
        lens = [len(o) for o in q["options"]]
        mx, mn = max(lens), min(lens)
        if lens[q["correct"]] == mx and lens.count(mx) == 1:
            longest += 1
            if lens[q["correct"]] >= 1.5 * sorted(lens, reverse=True)[1]:
                len15.append(q)
        if lens[q["correct"]] == mn and lens.count(mn) == 1:
            shortest += 1
            if lens[q["correct"]] <= 0.67 * sorted(lens)[1]:
                strong_shortest += 1
        for i, o in enumerate(q["options"]):
            (tok_c if i == q["correct"] else tok_w).update(set(re.findall(r"[a-ząćęłńóśźż]{3,}", o.lower())))
    token_ratios = sorted(((round(cc / max(tok_w.get(t, 0) / 3, 0.5), 1), cc, tok_w.get(t, 0), t) for t, cc in tok_c.items() if cc >= 20 and cc / max(tok_w.get(t, 0) / 3, 0.5) >= 2), reverse=True)
    tells = {"conj_only": only_correct(qs, CONJ), "paren_only": only_correct(qs, PAREN), "comma_only": only_correct(qs, COMMA), "catg_wrong": catg_wrong, "len15": len15}
    union = sorted({(q["sheet"], q["row"]) for lst in tells.values() for q in lst})
    report = {"n": n, "image": {"total": len(img_qs), "parenthetical": par, "embedded": len(img_qs) - par, "questions": img_qs}, "tells": tells, "union": [list(u) for u in union],
              "length": {"longest": longest, "shortest": shortest, "strong_shortest": strong_shortest},
              "token_ratios": [{"token": t, "ratio": r, "correct": cc, "wrong": wc} for r, cc, wc, t in token_ratios]}
    Path(sys.argv[2]).write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"pytań: {n} | image: {len(img_qs)} (nawiasowe {par}, wplecione {len(img_qs) - par})")
    for k, v in tells.items():
        print(f"{k}: {len(v)}")
    print(f"poprawna najdłuższa: {longest} ({100 * longest / n:.1f}%) | "
          f"najkrótsza: {shortest} ({100 * shortest / n:.1f}%), w tym silnie: {strong_shortest} ({100 * strong_shortest / n:.1f}%) | "
          f"unia: {len(union)} ({100 * len(union) / n:.1f}%)")
    print("tokeny ratio>=2 (top 5): " + ", ".join(f"{t} {r}" for r, cc, wc, t in token_ratios[:5]))


if __name__ == "__main__":
    main()
