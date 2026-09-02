#!/usr/bin/env python3
"""Nanosi propozycje zmian na xlsx i dopisuje je do dziennika. Jedyna droga zapisu do neuro_questions.xlsx w przebiegach detells."""
import argparse
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
from generate_quiz_data import cell_text, cites_letter  # noqa: E402

FIELD_COL = {"ID": 1, "Pytanie": 2, "A": 3, "B": 4, "C": 5, "D": 6, "Poprawna": 7, "Wyjaśnienie": 8, "Wykluczone": 10}
IMG = re.compile(r"\bimage\b", re.IGNORECASE)


def fail(msg):
    sys.exit(f"BŁĄD: {msg}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("proposals")
    ap.add_argument("--log", default=str(HERE.parent.parent / "docs" / "validation" / "applied-changes.json"))
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    import openpyxl
    props = json.loads(Path(args.proposals).read_text(encoding="utf-8"))
    if not isinstance(props, list) or not props:
        fail("plik propozycji pusty albo nie jest listą")
    wb = openpyxl.load_workbook(args.xlsx)
    log_path = None
    if args.apply:
        log_path = Path(args.log)
        if not log_path.exists():
            fail(f"plik dziennika nie istnieje: {log_path}")
    applied, skipped = [], 0
    for i, p in enumerate(props):
        missing = [k for k in ("sheet", "row", "id", "field", "from", "to", "why", "signals") if k not in p]
        if missing:
            fail(f"propozycja {i}: brak pól {missing}")
        if p["field"] not in FIELD_COL:
            fail(f"propozycja {i}: nieznane pole '{p['field']}'")
        if str(p["from"]) == str(p["to"]):
            fail(f"propozycja {i} ({p['sheet']} w.{p['row']}): 'from' == 'to' (pusta zmiana zaśmieca dziennik)")
        if p["sheet"] not in wb.sheetnames:
            fail(f"propozycja {i}: arkusz '{p['sheet']}' nie istnieje w skoroszycie")
        ws = wb[p["sheet"]]
        cur_id = cell_text(ws.cell(p["row"], FIELD_COL["ID"]))
        if cur_id != p["id"]:
            fail(f"{p['sheet']} w.{p['row']}: ID w arkuszu '{cur_id}' != '{p['id']}' (przesunięcie wierszy?)")
        cur = cell_text(ws.cell(p["row"], FIELD_COL[p["field"]]))
        if cur == p["to"]:
            skipped += 1  # idempotentny re-run: zmiana już w pliku (crash między save a logiem albo powtórka)
            print(f"POMINIĘTO (już naniesione) {p['sheet']} w.{p['row']} [{p['field']}]")
            continue
        if cur != p["from"]:
            fail(f"{p['sheet']} w.{p['row']} [{p['field']}]: wartość bieżąca różna od 'from' (stale?)\n  jest: {cur[:120]}\n  from: {str(p['from'])[:120]}")
        if not str(p["to"]).strip():
            fail(f"{p['sheet']} w.{p['row']} [{p['field']}]: puste 'to'")
        if p["field"] == "Wyjaśnienie":
            if IMG.search(p["to"]):
                fail(f"{p['sheet']} w.{p['row']}: 'to' nadal zawiera odniesienie do Image")
            if cites_letter(p["to"]):
                fail(f"{p['sheet']} w.{p['row']}: 'to' cytuje literę odpowiedzi (opcje są tasowane)")
        if args.apply:
            ws.cell(p["row"], FIELD_COL[p["field"]]).value = p["to"]
        applied.append(p)
        print(f"OK {p['sheet']} w.{p['row']} [{p['field']}]: {str(p['from'])[:60]!r} -> {str(p['to'])[:60]!r}")
    by_row = {}  # opcje wiersza po edycji muszą być parami różne; w dry-run nakładamy propozycje na odczyt
    for p in applied:
        if p["field"] in ("A", "B", "C", "D"):
            by_row.setdefault((p["sheet"], p["row"]), {})[p["field"]] = p["to"]
    for (sheet, row), edits in sorted(by_row.items()):
        ws = wb[sheet]
        opts = [re.sub(r"\s+", " ", str(edits.get(f, cell_text(ws.cell(row, FIELD_COL[f])))).lower()).strip() for f in ("A", "B", "C", "D")]
        if len(set(opts)) != 4:
            fail(f"{sheet} w.{row}: opcje po edycji nie są parami różne: {opts}")
    if args.apply:
        wb.save(args.xlsx)
        log = json.loads(log_path.read_text(encoding="utf-8"))
        seen = {json.dumps(e, sort_keys=True, ensure_ascii=False) for e in log}
        new = [p for p in applied if json.dumps(p, sort_keys=True, ensure_ascii=False) not in seen]
        log.extend(new)
        log_path.write_text(json.dumps(log, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"Zapisano {len(applied)} zmian do {args.xlsx} (nowe wpisy w dzienniku: {len(new)}, pominięte już naniesione: {skipped})")
    else:
        print(f"DRY-RUN: {len(applied)} propozycji poprawnych (pominięte już naniesione: {skipped}), nic nie zapisano")


if __name__ == "__main__":
    main()
