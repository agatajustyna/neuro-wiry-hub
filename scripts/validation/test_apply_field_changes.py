import json
import subprocess
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
FIELD_COL = {"ID": 1, "Pytanie": 2, "A": 3, "B": 4, "C": 5, "D": 6, "Poprawna": 7, "Wyjaśnienie": 8, "Wykluczone": 10}


def make_xlsx(path):
    wb = openpyxl.Workbook()
    wb.active.title = "Struktura"
    ws = wb.create_sheet("1.Testowy")
    ws.append(["ID", "Pytanie", "A", "B", "C", "D", "Poprawna", "Wyjaśnienie", "Tagi", "Wykluczone"])
    ws.append(["T-1", "Co?", "a1", "b1", "c1", "d1", "A", "Bo tak (Image 5).", "", ""])
    wb.save(path)


def run(xlsx, proposals, log, apply=False):
    cmd = [sys.executable, str(HERE / "apply_field_changes.py"), str(xlsx), str(proposals), "--log", str(log)]
    if apply:
        cmd.append("--apply")
    return subprocess.run(cmd, capture_output=True, text=True)


def test_apply_and_log(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    make_xlsx(xlsx)
    log = tmp_path / "log.json"
    log.write_text("[]", encoding="utf-8")
    props = [{"sheet": "1.Testowy", "row": 2, "id": "T-1", "field": "Wyjaśnienie", "from": "Bo tak (Image 5).", "to": "Bo tak.", "why": "usunięcie odniesienia do ryciny", "signals": "test"}]
    p = tmp_path / "p.json"
    p.write_text(json.dumps(props, ensure_ascii=False), encoding="utf-8")
    r = run(xlsx, p, log, apply=True)
    assert r.returncode == 0, r.stderr
    ws = openpyxl.load_workbook(xlsx)["1.Testowy"]
    assert ws.cell(2, FIELD_COL["Wyjaśnienie"]).value == "Bo tak."
    assert len(json.loads(log.read_text(encoding="utf-8"))) == 1


def test_stale_from_fails(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    make_xlsx(xlsx)
    log = tmp_path / "log.json"
    log.write_text("[]", encoding="utf-8")
    props = [{"sheet": "1.Testowy", "row": 2, "id": "T-1", "field": "Wyjaśnienie", "from": "INNA TREŚĆ", "to": "X.", "why": "w", "signals": "test"}]
    p = tmp_path / "p.json"
    p.write_text(json.dumps(props, ensure_ascii=False), encoding="utf-8")
    r = run(xlsx, p, log, apply=True)
    assert r.returncode != 0
    assert "INNA TREŚĆ"[:4] in r.stderr or "from" in r.stderr


def test_id_mismatch_fails(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    make_xlsx(xlsx)
    log = tmp_path / "log.json"
    log.write_text("[]", encoding="utf-8")
    props = [{"sheet": "1.Testowy", "row": 2, "id": "T-99", "field": "A", "from": "a1", "to": "a2", "why": "w", "signals": "test"}]
    p = tmp_path / "p.json"
    p.write_text(json.dumps(props, ensure_ascii=False), encoding="utf-8")
    r = run(xlsx, p, log, apply=True)
    assert r.returncode != 0


def test_letter_citation_in_explanation_fails(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    make_xlsx(xlsx)
    log = tmp_path / "log.json"
    log.write_text("[]", encoding="utf-8")
    props = [{"sheet": "1.Testowy", "row": 2, "id": "T-1", "field": "Wyjaśnienie", "from": "Bo tak (Image 5).", "to": "Odpowiedź A jest poprawna.", "why": "w", "signals": "test"}]
    p = tmp_path / "p.json"
    p.write_text(json.dumps(props, ensure_ascii=False), encoding="utf-8")
    r = run(xlsx, p, log, apply=True)
    assert r.returncode != 0


def test_rerun_is_idempotent(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    make_xlsx(xlsx)
    log = tmp_path / "log.json"
    log.write_text("[]", encoding="utf-8")
    props = [{"sheet": "1.Testowy", "row": 2, "id": "T-1", "field": "Wyjaśnienie", "from": "Bo tak (Image 5).", "to": "Bo tak.", "why": "w", "signals": "test"}]
    p = tmp_path / "p.json"
    p.write_text(json.dumps(props, ensure_ascii=False), encoding="utf-8")
    assert run(xlsx, p, log, apply=True).returncode == 0
    r2 = run(xlsx, p, log, apply=True)  # powtórka: crash-recovery / omyłkowe drugie uruchomienie
    assert r2.returncode == 0, r2.stderr
    assert "już naniesione" in r2.stdout
    assert len(json.loads(log.read_text(encoding="utf-8"))) == 1  # dziennik bez duplikatu


def test_duplicate_options_fail(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    make_xlsx(xlsx)
    log = tmp_path / "log.json"
    log.write_text("[]", encoding="utf-8")
    props = [{"sheet": "1.Testowy", "row": 2, "id": "T-1", "field": "B", "from": "b1", "to": "A1 ", "why": "w", "signals": "test"}]  # po normalizacji zrówna się z opcją A
    p = tmp_path / "p.json"
    p.write_text(json.dumps(props, ensure_ascii=False), encoding="utf-8")
    r = run(xlsx, p, log, apply=True)
    assert r.returncode != 0
    ws = openpyxl.load_workbook(xlsx)["1.Testowy"]
    assert ws.cell(2, FIELD_COL["B"]).value == "b1"  # nic nie zapisano


def test_dry_run_writes_nothing(tmp_path):
    xlsx = tmp_path / "t.xlsx"
    make_xlsx(xlsx)
    log = tmp_path / "log.json"
    log.write_text("[]", encoding="utf-8")
    props = [{"sheet": "1.Testowy", "row": 2, "id": "T-1", "field": "Wyjaśnienie", "from": "Bo tak (Image 5).", "to": "Bo tak.", "why": "w", "signals": "test"}]
    p = tmp_path / "p.json"
    p.write_text(json.dumps(props, ensure_ascii=False), encoding="utf-8")
    r = run(xlsx, p, log, apply=False)
    assert r.returncode == 0, r.stderr
    assert "DRY-RUN" in r.stdout
    ws = openpyxl.load_workbook(xlsx)["1.Testowy"]
    assert ws.cell(2, FIELD_COL["Wyjaśnienie"]).value == "Bo tak (Image 5)."  # xlsx nietknięty
    assert json.loads(log.read_text(encoding="utf-8")) == []  # dziennik nietknięty
